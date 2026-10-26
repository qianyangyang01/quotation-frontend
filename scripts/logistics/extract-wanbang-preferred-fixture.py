"""Read the supplied 2026-10-26 workbook without changing it; record source cells for regression tests."""
import hashlib
import json
import pathlib
import sys
import openpyxl

source = pathlib.Path(sys.argv[1])
target = pathlib.Path(sys.argv[2])
book = openpyxl.load_workbook(source, read_only=True, data_only=True)
countries = {"德国": "DE", "法国": "FR", "意大利": "IT", "西班牙": "ES", "葡萄牙": "PT", "荷兰": "NL"}
channels = []
for name in ["万邦优选专线普货", "万邦优选专线含电", "万邦优选专线敏感"]:
    sheet = book[name]
    rows = []
    labels = {}
    for number, values in enumerate(sheet.iter_rows(min_row=6, max_row=32, max_col=11, values_only=True), 6):
        for column in range(1, 7):
            if values[column] is not None:
                labels[column] = values[column]
        assert values[7] and isinstance(values[8], (int, float)) and isinstance(values[9], (int, float))
        rows.append(dict(sourceRow=number, country=labels[1], countryCode=countries[labels[1]], product=labels[2],
                         code=labels[3], eta=labels[4], limit=labels[5], size=labels[6], weight=values[7],
                         pricePerKg=values[8], registrationFee=values[9]))
    channels.append(dict(channelName=name, title=sheet.cell(1, 4).value, rows=rows, taxNote=sheet.cell(43, 2).value))
target.parent.mkdir(parents=True, exist_ok=True)
target.write_text(json.dumps(dict(source=source.name, sha256=hashlib.sha256(source.read_bytes()).hexdigest(), channels=channels), ensure_ascii=False, indent=2), encoding="utf-8")
print(f"Extracted {len(channels)} channels, {sum(len(c['rows']) for c in channels)} rows")
