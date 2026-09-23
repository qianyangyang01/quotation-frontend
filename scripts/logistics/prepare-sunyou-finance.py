"""Read the two approved workbooks; prepare only the three Sunyou settings.

No network/database writes. Source workbooks are never changed. The finance
snapshot and generated evidence stay outside Git in the release directory.
"""
import copy
import datetime
import hashlib
import json
import pathlib
import sys
from decimal import Decimal, ROUND_HALF_UP

import openpyxl

TARGETS = {
    '顺邮宝挂号': (643, 'SYBRAM', ['普货', '化妆品', '服装', '带电']),
    '顺速宝(特货)': (644, 'SSBRAM', ['化妆品', '带电']),
    '顺速宝Plus': (645, 'SYEPL', ['化妆品', '带电', '纯电', '香水']),
}
TAX_COLUMNS = {
    '欧盟关税': ('欧盟', 'channelRules'), '美国关税': ('美国', 'channelRules'),
    '新西兰': ('新西兰', 'channelRules'), '墨西哥关税': ('墨西哥', 'channelRules'),
    '阿拉伯联合酋长国关税': ('阿拉伯联合酋长国', 'channelRules'),
    '沙特阿拉伯关税': ('沙特阿拉伯', 'channelRules'), '哥伦比亚': ('哥伦比亚', 'channelRules'),
    '约旦关税': ('约旦', 'channelRules'), '摩洛哥-增值税': ('摩洛哥', 'channelRules'),
    '阿曼': ('阿曼', 'channelRules'), '卡塔尔-处理费': ('卡塔尔', 'channelRules'),
    '罗马尼亚-关税处理费': ('罗马尼亚', 'handlingRules'),
}
REMOTE_COLUMNS = {'新西兰-偏远', '英国-偏远', '加拿大-偏远'}
BLOCKED_COUNTRIES = ['US', 'GB', 'CA', 'AU', 'NZ', 'MX', 'RO']


def load_source(path):
    # WPS files can incorrectly advertise A1:A1 as their worksheet dimension.
    # Normal loading reads the actual cell records, including the appended rows.
    formulas = openpyxl.load_workbook(path, data_only=False, read_only=False)
    cached = openpyxl.load_workbook(path, data_only=True, read_only=False)
    rows = {}
    for sheet in formulas:
        headers = [c.value for c in sheet[1]]
        for row in sheet.iter_rows(min_row=2):
            if row[0].value != '顺友':
                continue
            name = row[1].value
            assert name in TARGETS and name not in rows, f'Unknown/duplicate Sunyou channel: {name}'
            cells = []
            for cell in row[2:]:
                if cell.column > len(headers) or headers[cell.column - 1] is None:
                    assert cell.value is None
                    continue
                cells.append({'column': headers[cell.column - 1], 'cell': cell.coordinate,
                              'formulaOrValue': cell.value, 'value': cached[sheet.title][cell.coordinate].value})
            rows[name] = {'sheet': sheet.title, 'row': row[0].row, 'cells': cells}
    assert set(rows) == set(TARGETS)
    return {'file': pathlib.Path(path).name, 'sha256': hashlib.sha256(pathlib.Path(path).read_bytes()).hexdigest(), 'rows': rows}


def prepare(snapshot, attributes, taxes, remote_scope, uk_surcharge):
    assert remote_scope == 'whole-country'
    assert uk_surcharge == 'not-applicable'
    assert not snapshot['paused'] and snapshot['activeImports'] == 0
    channels = [c for c in snapshot['channels'] if c['provider'] == '顺友']
    assert {c['name'] for c in channels} == set(TARGETS)
    keys = {c['name']: f"{c['ruleId']}::顺友::{c['code']}" for c in channels}
    for c in channels:
        rule_id, product, expected = TARGETS[c['name']]
        assert c['ruleId'] == rule_id and c['enabled'] and c['quoteReady'] and c['company']['enabled']
        assert c['company']['productCodes'] == [product]
        cells = attributes['rows'][c['name']]['cells']
        assert len(cells) == 9 and all(x['value'] in (None, '√') for x in cells)
        assert [x['column'] for x in cells if x['value'] == '√'] == expected
    before = {k: snapshot['settings'][k]['payload'] for k in ('tax-settings', 'channel-policies', 'surcharge-settings')}
    after = copy.deepcopy(before)
    stamp = datetime.datetime.now(datetime.timezone.utc).isoformat()
    tax_rows = {r['country']: r for r in after['tax-settings']['countries']}
    changes = []
    remote = []
    for name, row in taxes['rows'].items():
        assert {c['column'] for c in row['cells']} == set(TAX_COLUMNS) | REMOTE_COLUMNS
        for cell in row['cells']:
            column, value = cell['column'], cell['value']
            source = f"{taxes['file']} / {row['sheet']}!{cell['cell']}"
            if column in REMOTE_COLUMNS:
                assert value == '不发'
                remote.append({'channel': name, **cell, 'source': source, 'scope': remote_scope})
                continue
            country, field = TAX_COLUMNS[column]
            assert tax_rows[country]['selected']
            mode = {'不发': 'unavailable', '无': 'no-tax', '含': 'exempt'}.get(value)
            if mode is None:
                assert isinstance(value, (int, float)) and not isinstance(value, bool)
                amount = float(Decimal(str(value)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP))
                mode = 'fixed-order'
                assert column == '欧盟关税' and amount == 3.51 and cell['formulaOrValue'] == '=23.52/6.7'
            else:
                amount = 0
            rule = {'key': keys[name], 'mode': mode, 'amount': amount, 'perKg': 0, 'currency': 'USD', 'source': source}
            entries = tax_rows[country].setdefault(field, [])
            assert not any(r['key'] == keys[name] for r in entries), f'Existing target rule: {country}/{name}'
            entries.append(rule)
            changes.append({'channel': name, 'country': country, 'field': field, 'rule': rule, 'sourceCell': cell})
    assert len(changes) == 36 and len(remote) == 9
    if remote_scope == 'whole-country':
        for name in TARGETS:
            for country in ('英国', '加拿大', '澳大利亚'):
                assert tax_rows[country]['selected']
                rule = {'key': keys[name], 'mode': 'unavailable', 'amount': 0, 'perKg': 0, 'currency': 'USD', 'source': '用户确认美国、英国、加拿大、澳大利亚全境不发；奥已确认为澳大利亚'}
                assert not any(r['key'] == keys[name] for r in tax_rows[country].get('channelRules', []))
                tax_rows[country].setdefault('channelRules', []).append(rule)
                changes.append({'channel': name, 'country': country, 'field': 'channelRules', 'rule': rule})
    # Give only the marked attributes to these channels, for every permitted published
    # destination. Existing country enablement is deliberately left intact.
    country_settings = snapshot['settings']['country-classification']['payload']
    alias_rows = json.loads((pathlib.Path(__file__).resolve().parents[2] / 'backend/src/main/resources/country-aliases.json').read_text(encoding='utf-8'))
    aliases = {value.upper(): row[0] for row in alias_rows for value in row}
    def identity(value):
        return aliases.get(value.strip().upper(), value.strip().upper())
    unsupported_aliases = []
    authorizations = []
    for policy in after['channel-policies']:
        existing = {r['country']: r for r in policy['countryRules']}
        for channel in channels:
            name = channel['name']
            if policy['category'] not in TARGETS[name][2]:
                continue
            assert policy['enabled']
            for destination in channel['countries']:
                if destination['code'] in BLOCKED_COUNTRIES:
                    continue
                candidates = [r for r in country_settings if r['code'] == destination['code']]
                choices = [r for r in candidates if identity(r['country']) in (identity(destination['code']), identity(destination['name']))]
                for row in candidates:
                    if row not in choices and row not in unsupported_aliases:
                        unsupported_aliases.append(row)
                if not choices:
                    choices = [{'country': destination['name'], 'code': destination['code'], 'enabled': False,
                                'stage': 'rare', 'sortOrder': 10000, 'continent': ''}]
                for meta in choices:
                    country = meta['country']
                    if country not in existing:
                        entry = {k: meta[k] for k in ('country', 'stage', 'sortOrder', 'continent')}
                        entry.update(allowedChannels=[], unavailableChannels=[])
                        policy['countryRules'].append(entry)
                        existing[country] = entry
                    entry = existing[country]
                    if keys[name] not in entry['allowedChannels']:
                        entry['allowedChannels'].append(keys[name])
                        authorizations.append({'channel': name, 'attribute': policy['category'], 'country': country,
                                               'code': destination['code'], 'countryEnabled': meta['enabled']})
        if any(x['attribute'] == policy['category'] for x in authorizations):
            policy['updatedAt'] = stamp
    after['tax-settings']['updatedAt'] = stamp
    # No existing rule/authorization/setting may be overwritten or removed.
    for old in before['tax-settings']['countries']:
        new = tax_rows[old['country']]
        for field in ('channelRules', 'handlingRules'):
            assert new.get(field, [])[:len(old.get(field, []))] == old.get(field, [])
        assert {k: v for k, v in new.items() if k not in ('channelRules', 'handlingRules')} == {k: v for k, v in old.items() if k not in ('channelRules', 'handlingRules')}
    assert after['tax-settings']['providers'] == before['tax-settings']['providers']
    for old, new in zip(before['channel-policies'], after['channel-policies'], strict=True):
        new_rules = {r['country']: r for r in new['countryRules']}
        for rule in old['countryRules']:
            restored = copy.deepcopy(new_rules[rule['country']])
            restored['allowedChannels'] = [k for k in restored['allowedChannels'] if k not in keys.values()]
            assert restored == rule
        for rule in new['countryRules'][len(old['countryRules']):]:
            assert set(rule['allowedChannels']) <= set(keys.values())
    report = {'channels': [{'name': c['name'], 'key': keys[c['name']], 'id': c['id'], 'attributes': TARGETS[c['name']][2]} for c in channels],
              'sources': [attributes, taxes], 'taxChanges': changes, 'authorizations': authorizations, 'remoteRestrictions': remote,
              'remoteScope': remote_scope, 'ukSurcharge': uk_surcharge, 'blockedCountryCodes': BLOCKED_COUNTRIES, 'unrecognizedExistingCountryAliases': unsupported_aliases,
              'fromVersions': {k: snapshot['settings'][k]['version'] for k in after}}
    return {k: v for k, v in after.items() if v != before[k]}, report


if __name__ == '__main__':
    snapshot_path, attribute_path, tax_path, directory, scope, uk_surcharge = sys.argv[1:]
    snapshot = json.loads(pathlib.Path(snapshot_path).read_text(encoding='utf-8-sig'))
    prepared, report = prepare(snapshot, load_source(attribute_path), load_source(tax_path), scope, uk_surcharge)
    output = pathlib.Path(directory)
    output.mkdir(parents=True, exist_ok=True)
    for name, value in [('finance-prepared', prepared), ('change-report', report)]:
        (output / f'{name}.json').write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'taxRules': len(report['taxChanges']), 'attributeAuthorizations': len(report['authorizations']),
                      'sourceCells': sum(len(r['cells']) for s in report['sources'] for r in s['rows'].values()),
                      'remoteRestrictions': len(report['remoteRestrictions']), 'updatedSettings': list(prepared)}, ensure_ascii=False))
